#!/usr/bin/env python3
"""Keep only hospitals listed on EQA_v1.xlsx sheet 1 (EQA信息表) VISIBLE rows.

Excel AutoFilter hides rows with 26年EQA标签 = '/' (~1073). Canonical
universe = visible rows only (~364, same as 省考客户 scale).

Matching:
  - Prefer exact customer_code match to sheet1 客户编码.
  - Else leading-zero-insensitive match (化免短码 → 同一家).
  - One keep row per sheet1 code (prefer exact padded code).
  - Recode kept unpadded rows to the sheet1 padded code.
  - Delete all other hospitals (demo seeds, tests, short-code duplicates,
    and sheet1 rows that are filter-hidden / tag='/').

FK handling before DELETE:
  - survey_assignments: ON DELETE RESTRICT → delete assignments for drop set.
  - Other hospital FKs are CASCADE or SET NULL.

Usage:
  DATABASE_URL=postgres://... python3 scripts/prune_hospitals_to_eqa_sheet1.py \\
    [--xlsx scripts/data/EQA_v1.xlsx] [--dry-run|--commit]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as e:
    print("need openpyxl:", e, file=sys.stderr)
    sys.exit(2)

try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError as e:
    print("need psycopg:", e, file=sys.stderr)
    sys.exit(2)


SHEET1_NAME = "EQA信息表"


def norm_ws(s: Any) -> str:
    if s is None:
        return ""
    t = str(s).replace("\n", "").replace("\r", "").strip()
    return re.sub(r"\s+", " ", t)


def norm_code(s: Any) -> str:
    t = norm_ws(s)
    if not t:
        return ""
    if t.isdigit():
        stripped = t.lstrip("0")
        return stripped if stripped else "0"
    return t


EMPTY_TAGS = {"", "/", "-", "—", "#N/A", "N/A"}


def is_empty_eqa_tag(v: Any) -> bool:
    return norm_ws(v) in EMPTY_TAGS


def load_sheet1_codes(path: Path) -> list[str]:
    """Visible EQA信息表 rows only (skip AutoFilter-hidden / empty tag '/').

    openpyxl needs read_only=False for ws.row_dimensions[r].hidden.
    """
    # NOT read_only: row_dimensions.hidden is unavailable in read_only mode
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET1_NAME in wb.sheetnames:
        ws = wb[SHEET1_NAME]
    else:
        ws = wb[wb.sheetnames[0]]
    if ws.max_row is None or ws.max_row < 2:
        raise SystemExit(f"empty sheet in {path}")
    header = [norm_ws(ws.cell(1, c).value) for c in range(1, (ws.max_column or 0) + 1)]
    try:
        code_i = header.index("客户编码") + 1  # 1-based
    except ValueError:
        code_i = 3
    try:
        tag_i = header.index("26年EQA标签") + 1
    except ValueError:
        tag_i = 6
    codes: list[str] = []
    seen: set[str] = set()
    skipped_hidden = 0
    skipped_empty_tag = 0
    for r in range(2, ws.max_row + 1):
        if bool(ws.row_dimensions[r].hidden):
            skipped_hidden += 1
            continue
        if is_empty_eqa_tag(ws.cell(r, tag_i).value):
            skipped_empty_tag += 1
            continue
        code = norm_ws(ws.cell(r, code_i).value)
        if not code or code in {"#N/A", "N/A", "/", "-"}:
            continue
        if code in seen:
            continue
        seen.add(code)
        codes.append(code)
    print(
        f"sheet1 visible codes={len(codes)} "
        f"(skipped_hidden={skipped_hidden}, skipped_empty_tag={skipped_empty_tag})",
        flush=True,
    )
    return codes


def run(xlsx: Path, dsn: str, *, dry: bool, commit: bool) -> dict:
    sheet_codes = load_sheet1_codes(xlsx)
    stats: dict[str, Any] = {
        "xlsx": str(xlsx),
        "sheet1_unique_codes": len(sheet_codes),
        "before_hospitals": 0,
        "keep": 0,
        "recode": 0,
        "deleted": 0,
        "assignments_deleted": 0,
        "dry_run": dry,
        "committed": False,
        "sample_deleted": [],
        "sample_recoded": [],
    }

    with psycopg.connect(dsn, row_factory=dict_row) as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS n FROM hospitals")
            stats["before_hospitals"] = cur.fetchone()["n"]

            cur.execute("CREATE TEMP TABLE eqa_keep_codes (code text PRIMARY KEY)")
            cur.executemany(
                "INSERT INTO eqa_keep_codes (code) VALUES (%s)",
                [(c,) for c in sheet_codes],
            )

            # Prefer exact match; else leading-zero match; one row per sheet ncode.
            cur.execute(
                """
                CREATE TEMP TABLE sheet AS
                SELECT trim(code) AS raw,
                       CASE WHEN code ~ '^[0-9]+$' THEN NULLIF(ltrim(trim(code),'0'),'') ELSE trim(code) END AS ncode
                FROM eqa_keep_codes;

                CREATE TEMP TABLE hnorm AS
                SELECT id, name, trim(customer_code) AS code,
                       CASE WHEN customer_code ~ '^[0-9]+$' THEN NULLIF(ltrim(trim(customer_code),'0'),'') ELSE trim(customer_code) END AS ncode,
                       EXISTS (
                         SELECT 1 FROM eqa_keep_codes k WHERE trim(k.code)=trim(h.customer_code)
                       ) AS exact
                FROM hospitals h
                WHERE customer_code IS NOT NULL AND trim(customer_code) <> '';

                CREATE TEMP TABLE keep_final AS
                WITH ranked AS (
                  SELECT h.id, h.name, h.code, s.raw AS sheet_raw, h.exact,
                         ROW_NUMBER() OVER (
                           PARTITION BY s.ncode
                           ORDER BY h.exact DESC, (h.code = s.raw) DESC, h.id ASC
                         ) AS rn
                  FROM sheet s
                  JOIN hnorm h ON h.ncode = s.ncode
                )
                SELECT id, name, code, sheet_raw, exact FROM ranked WHERE rn = 1;

                CREATE TEMP TABLE drop_final AS
                SELECT id, name, customer_code, status
                FROM hospitals
                WHERE id NOT IN (SELECT id FROM keep_final);
                """
            )

            cur.execute("SELECT COUNT(*) AS n FROM keep_final")
            stats["keep"] = cur.fetchone()["n"]
            cur.execute("SELECT COUNT(*) AS n FROM drop_final")
            stats["deleted"] = cur.fetchone()["n"]
            cur.execute(
                "SELECT COUNT(*) AS n FROM keep_final WHERE NOT exact OR code IS DISTINCT FROM sheet_raw"
            )
            stats["recode"] = cur.fetchone()["n"]

            cur.execute(
                """
                SELECT id, left(name,40) AS name, coalesce(customer_code,'') AS code, status
                FROM drop_final ORDER BY id LIMIT 40
                """
            )
            stats["sample_deleted"] = [dict(r) for r in cur.fetchall()]

            cur.execute(
                """
                SELECT id, left(name,40) AS name, code AS old_code, sheet_raw AS new_code
                FROM keep_final
                WHERE code IS DISTINCT FROM sheet_raw
                ORDER BY id LIMIT 20
                """
            )
            stats["sample_recoded"] = [dict(r) for r in cur.fetchall()]

            # Sanity: keep count must equal sheet1 codes (every sheet code matched)
            cur.execute(
                """
                SELECT COUNT(*) AS missing
                FROM sheet s
                WHERE NOT EXISTS (SELECT 1 FROM keep_final k WHERE k.sheet_raw = s.raw)
                """
            )
            missing = cur.fetchone()["missing"]
            stats["sheet_codes_unmatched"] = missing
            if missing:
                cur.execute(
                    """
                    SELECT s.raw FROM sheet s
                    WHERE NOT EXISTS (SELECT 1 FROM keep_final k WHERE k.sheet_raw = s.raw)
                    LIMIT 20
                    """
                )
                stats["unmatched_sample"] = [r["raw"] for r in cur.fetchall()]
                conn.rollback()
                print(json.dumps(stats, ensure_ascii=False, indent=2, default=str))
                raise SystemExit(
                    f"ABORT: {missing} sheet1 codes have no hospital match; refusing prune"
                )

            if stats["keep"] != len(sheet_codes):
                conn.rollback()
                print(json.dumps(stats, ensure_ascii=False, indent=2, default=str))
                raise SystemExit(
                    f"ABORT: keep={stats['keep']} != sheet1={len(sheet_codes)}"
                )

            # Apply mutations
            cur.execute(
                """
                DELETE FROM survey_assignments
                WHERE hospital_id IN (SELECT id FROM drop_final)
                """
            )
            stats["assignments_deleted"] = cur.rowcount

            cur.execute(
                """
                UPDATE hospitals h
                SET customer_code = k.sheet_raw, updated_at = NOW()
                FROM keep_final k
                WHERE h.id = k.id AND h.customer_code IS DISTINCT FROM k.sheet_raw
                """
            )
            stats["recode_applied"] = cur.rowcount

            cur.execute("DELETE FROM hospitals WHERE id IN (SELECT id FROM drop_final)")
            stats["deleted_applied"] = cur.rowcount

            cur.execute("SELECT COUNT(*) AS n FROM hospitals")
            stats["after_hospitals"] = cur.fetchone()["n"]

            if dry or not commit:
                conn.rollback()
                stats["committed"] = False
                stats["note"] = "rolled back (dry-run or --commit not set)"
            else:
                conn.commit()
                stats["committed"] = True

    return stats


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    root = Path(__file__).resolve().parents[1]
    ap.add_argument("--xlsx", type=Path, default=root / "scripts/data/EQA_v1.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--commit", action="store_true")
    ap.add_argument("--dsn", default=os.environ.get("DATABASE_URL", ""))
    args = ap.parse_args()
    if not args.dsn:
        print("DATABASE_URL or --dsn required", file=sys.stderr)
        sys.exit(2)
    if not args.xlsx.exists():
        print("xlsx missing:", args.xlsx, file=sys.stderr)
        sys.exit(2)
    commit = bool(args.commit) and not args.dry_run
    stats = run(args.xlsx, args.dsn, dry=args.dry_run, commit=commit)
    print(json.dumps(stats, ensure_ascii=False, indent=2, default=str))
    if not stats.get("committed") and not args.dry_run and not args.commit:
        print("Parsed OK but --commit not set; rolled back", file=sys.stderr)


if __name__ == "__main__":
    main()
