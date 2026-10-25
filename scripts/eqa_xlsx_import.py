#!/usr/bin/env python3
"""Idempotent EQA xlsx → Postgres import (conflict-safe).

Usage:
  DATABASE_URL=postgres://... python3 scripts/eqa_xlsx_import.py [--xlsx PATH] [--dry-run] [--commit]

Rules (mandatory):
  - Hospital universe = sheet 「EQA信息表」 only (国考/省考 never add hospitals).
  - Match by normalized customer_code (strip, leading-zero-insensitive); store padded form.
  - Name fallback only if code empty on BOTH sides.
  - INSERT new hospitals only; on match UPDATE empty columns only; never overwrite
    non-empty name/province/city/level/type/remark/status/attributes.
  - Disabled names → status=inactive on INSERT only (do not flip existing status).
  - NEW hospital defaults: province=安徽省, city from 分公司 or 合肥市, level=二级甲等, type=综合医院.
  - dual_major: 是/是双大/1 → attr if missing; #N/A/空/斜杠 skip.
  - EQA tags year=2026 on eqa_participation.tags; national preferred.
  - National 1 → items true; skip 0s; existing items never deleted/updated.
  - Does not touch field_registry / assay_projects / surveys.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

try:
    import openpyxl
except ImportError as e:
    print("need openpyxl:", e, file=sys.stderr)
    sys.exit(2)

try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError as e:
    print("need psycopg:", e, file=sys.stderr)
    sys.exit(2)


YEAR = 2026
META_COLS = {"区域", "分公司", "客户编码", "客户名称", "责任人", "26年EQA标签", "26年是否双大"}


def norm_ws(s: Any) -> str:
    if s is None:
        return ""
    t = str(s).replace("\n", "").replace("\r", "").strip()
    return re.sub(r"\s+", " ", t)


def norm_label(s: str) -> str:
    return re.sub(r"\s+", "", norm_ws(s))


def norm_code(s: Any) -> str:
    """Compare key: trim + strip leading zeros (keep '0' if all zeros)."""
    t = norm_ws(s)
    if not t:
        return ""
    if t.isdigit():
        stripped = t.lstrip("0")
        return stripped if stripped else "0"
    return t


def store_code(s: Any) -> str:
    """Prefer original padded digits as stored form."""
    t = norm_ws(s)
    return t


def is_disabled_name(name: str) -> bool:
    return "禁用" in name


def parse_dual_major(v: Any) -> Optional[str]:
    t = norm_ws(v)
    if not t or t in {"#N/A", "N/A", "/", "-", "—", "无", "否", "0"}:
        return None
    if t in {"是", "是双大", "1", "Y", "y", "true", "TRUE"}:
        return "是"
    if "双大" in t:
        return "是"
    return None


def parse_eqa_tag(v: Any) -> Optional[str]:
    t = norm_ws(v)
    if not t or t in {"/", "-", "—", "#N/A"}:
        return None
    return t


def city_from_branch(branch: str) -> str:
    b = norm_ws(branch)
    if not b:
        return "合肥市"
    if b.endswith("分公司"):
        b = b[: -len("分公司")]
    b = b.strip()
    if not b:
        return "合肥市"
    if b.endswith(("市", "地区", "州", "县")):
        return b
    return b + "市"


@dataclass
class Stats:
    hospitals_inserted: int = 0
    hospitals_matched: int = 0
    hospitals_filled: int = 0
    hospitals_skipped: int = 0
    dual_major_written: int = 0
    dual_major_skipped: int = 0
    participations_created: int = 0
    participations_existing: int = 0
    participations_tags_filled: int = 0
    items_inserted: int = 0
    items_skipped_existing: int = 0
    subject_unmapped: set[str] = field(default_factory=set)
    errors: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "hospitals_inserted": self.hospitals_inserted,
            "hospitals_matched": self.hospitals_matched,
            "hospitals_filled": self.hospitals_filled,
            "hospitals_skipped": self.hospitals_skipped,
            "dual_major_written": self.dual_major_written,
            "dual_major_skipped": self.dual_major_skipped,
            "participations_created": self.participations_created,
            "participations_existing": self.participations_existing,
            "participations_tags_filled": self.participations_tags_filled,
            "items_inserted": self.items_inserted,
            "items_skipped_existing": self.items_skipped_existing,
            "subject_unmapped": sorted(self.subject_unmapped),
            "errors": self.errors[:50],
            "error_count": len(self.errors),
        }


def sheet_headers(ws) -> dict[str, int]:
    out = {}
    for c in range(1, ws.max_column + 1):
        h = norm_ws(ws.cell(1, c).value)
        if h:
            out[h] = c
    return out


def subject_headers(headers: dict[str, int]) -> list[str]:
    return [h for h in headers if h not in META_COLS]


def load_xlsx(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    info = wb["EQA信息表"]
    guo = wb["国考客户"]
    # 省考客户 subjects empty — we may still read tags/names but skip subject facts
    sheng = wb["省考客户"]

    info_h = sheet_headers(info)
    guo_h = sheet_headers(guo)
    sheng_h = sheet_headers(sheng)

    hospitals = {}  # norm_code -> record (info sheet is universe)

    def read_rows(ws, headers, with_owner=False):
        rows = []
        code_c = headers.get("客户编码")
        name_c = headers.get("客户名称")
        region_c = headers.get("区域")
        branch_c = headers.get("分公司")
        dual_c = headers.get("26年是否双大")
        tag_c = headers.get("26年EQA标签")
        owner_c = headers.get("责任人") if with_owner else None
        subj = subject_headers(headers)
        for r in range(2, ws.max_row + 1):
            code = store_code(ws.cell(r, code_c).value) if code_c else ""
            name = norm_ws(ws.cell(r, name_c).value) if name_c else ""
            if not code and not name:
                continue
            row = {
                "code": code,
                "name": name,
                "region": norm_ws(ws.cell(r, region_c).value) if region_c else "",
                "branch": norm_ws(ws.cell(r, branch_c).value) if branch_c else "",
                "dual_raw": ws.cell(r, dual_c).value if dual_c else None,
                "tag": parse_eqa_tag(ws.cell(r, tag_c).value) if tag_c else None,
                "owner": norm_ws(ws.cell(r, owner_c).value) if owner_c else "",
                "subjects": {},
            }
            for s in subj:
                v = ws.cell(r, headers[s]).value
                row["subjects"][s] = v
            rows.append(row)
        return rows, subj

    info_rows, info_subj = read_rows(info, info_h)
    guo_rows, guo_subj = read_rows(guo, guo_h, with_owner=True)
    sheng_rows, sheng_subj = read_rows(sheng, sheng_h)

    for row in info_rows:
        key = norm_code(row["code"]) or f"name:{row['name']}"
        hospitals[key] = row

    # Overlay tags/owner from 国考/省考 onto sheet1 hospitals ONLY.
    # Never insert hospitals that are not on 「EQA信息表」 (canonical universe).
    guo_by_key = {}
    for row in guo_rows:
        key = norm_code(row["code"]) or f"name:{row['name']}"
        guo_by_key[key] = row
        if key not in hospitals:
            continue
        h = hospitals[key]
        if not h.get("tag") and row.get("tag"):
            h["tag"] = row["tag"]
        if not h.get("owner") and row.get("owner"):
            h["owner"] = row["owner"]

    for row in sheng_rows:
        key = norm_code(row["code"]) or f"name:{row['name']}"
        if key not in hospitals:
            continue
        h = hospitals[key]
        if not h.get("tag") and row.get("tag"):
            h["tag"] = row["tag"]

    return {
        "hospitals": hospitals,
        "guo_by_key": guo_by_key,
        "guo_subj": guo_subj,
        "info_subj": info_subj,
        "sheng_subj": sheng_subj,
    }


def load_field_maps(conn) -> dict[str, str]:
    """normalized label -> field_code for national boolean subjects."""
    m = {}
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT code, name FROM field_registry
            WHERE parent_code = 'scope.national' AND value_type = 'boolean'
            """
        )
        for code, name in cur.fetchall():
            m[norm_label(name)] = code
    return m


def load_db_hospitals(conn) -> tuple[dict[str, dict], dict[str, dict]]:
    """Returns (by_norm_code, by_exact_name) where name map only used when codes empty."""
    by_code: dict[str, dict] = {}
    by_name: dict[str, list[dict]] = defaultdict(list)
    with conn.cursor(row_factory=dict_row) as cur:
        cur.execute(
            """
            SELECT id, name, province, city, district, level, type, status,
                   address, remark,
                   COALESCE(customer_code,'') AS customer_code,
                   COALESCE(region,'') AS region,
                   COALESCE(branch_office,'') AS branch_office
            FROM hospitals
            """
        )
        rows = list(cur.fetchall())
        cur.execute(
            """
            SELECT hospital_id, attr_key, attr_value FROM hospital_attributes
            WHERE attr_key = 'dual_major'
            """
        )
        dual = {r["hospital_id"]: r["attr_value"] for r in cur.fetchall()}
    for r in rows:
        r = dict(r)
        r["dual_major"] = dual.get(r["id"])
        nc = norm_code(r["customer_code"])
        if nc:
            # first wins; keep existing
            by_code.setdefault(nc, r)
        by_name[r["name"]].append(r)
    return by_code, by_name


def match_hospital(row, by_code, by_name) -> Optional[dict]:
    nc = norm_code(row["code"])
    if nc and nc in by_code:
        return by_code[nc]
    if not nc:
        # fallback: exact name only if DB row also has empty code
        cands = [h for h in by_name.get(row["name"], []) if not norm_code(h["customer_code"])]
        if len(cands) == 1:
            return cands[0]
    return None


def fill_empty_hospital(conn, hid: int, row: dict, existing: dict, stats: Stats, dry: bool) -> None:
    sets = []
    args = []
    # only empty columns: customer_code, region, branch_office
    if not norm_ws(existing.get("customer_code")) and row["code"]:
        sets.append(f"customer_code = %s")
        args.append(store_code(row["code"]))
    if not norm_ws(existing.get("region")) and row.get("region"):
        sets.append("region = %s")
        args.append(row["region"])
    if not norm_ws(existing.get("branch_office")) and row.get("branch"):
        sets.append("branch_office = %s")
        args.append(row["branch"])
    if not sets:
        return
    args.append(hid)
    sql = f"UPDATE hospitals SET {', '.join(sets)}, updated_at = NOW() WHERE id = %s"
    if dry:
        stats.hospitals_filled += 1
        return
    with conn.cursor() as cur:
        cur.execute(sql, args)
    stats.hospitals_filled += 1


def insert_hospital(conn, row: dict, stats: Stats, dry: bool) -> Optional[int]:
    name = row["name"]
    status = "inactive" if is_disabled_name(name) else "active"
    province = "安徽省"
    city = city_from_branch(row.get("branch") or "")
    level = "二级甲等"
    typ = "综合医院"
    code = store_code(row["code"])
    region = row.get("region") or ""
    branch = row.get("branch") or ""
    if dry:
        stats.hospitals_inserted += 1
        return -1  # sentinel for dry-run chain
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO hospitals (
                    name, province, city, district, level, type, status,
                    address, remark, customer_code, region, branch_office, archive
                ) VALUES (
                    %s, %s, %s, '', %s, %s, %s,
                    '', '', %s, %s, %s, '{}'::jsonb
                )
                RETURNING id
                """,
                (name, province, city, level, typ, status, code, region, branch),
            )
            hid = cur.fetchone()[0]
        stats.hospitals_inserted += 1
        return hid
    except Exception as e:
        # unique (name, province, city) collision → treat as match if found
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(
                """
                SELECT id, name, province, city, district, level, type, status,
                       address, remark,
                       COALESCE(customer_code,'') AS customer_code,
                       COALESCE(region,'') AS region,
                       COALESCE(branch_office,'') AS branch_office
                FROM hospitals
                WHERE name = %s AND province = %s AND city = %s
                LIMIT 1
                """,
                (name, province, city),
            )
            existing = cur.fetchone()
        if existing:
            stats.hospitals_matched += 1
            fill_empty_hospital(conn, existing["id"], row, existing, stats, dry=False)
            return existing["id"]
        stats.errors.append(f"insert hospital fail {name}/{code}: {e}")
        return None


def ensure_dual_major(conn, hid: int, dual: Optional[str], existing_attr: Optional[str], stats: Stats, dry: bool):
    if not dual:
        stats.dual_major_skipped += 1
        return
    if existing_attr:
        stats.dual_major_skipped += 1
        return
    if dry:
        stats.dual_major_written += 1
        return
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO hospital_attributes (hospital_id, attr_key, attr_value)
            VALUES (%s, 'dual_major', %s)
            ON CONFLICT (hospital_id, attr_key) DO NOTHING
            """,
            (hid, dual),
        )
        if cur.rowcount:
            stats.dual_major_written += 1
        else:
            stats.dual_major_skipped += 1


def ensure_national_participation(
    conn,
    hid: int,
    tag: Optional[str],
    ones: list[str],
    stats: Stats,
    dry: bool,
):
    """Create or reuse (hospital,2026,national); add missing true items; fill empty tags."""
    if dry:
        # approximate
        if ones or tag:
            stats.participations_created += 1
            stats.items_inserted += len(ones)
            if tag:
                stats.participations_tags_filled += 1
        return

    with conn.cursor(row_factory=dict_row) as cur:
        cur.execute(
            """
            SELECT id, tags FROM eqa_participation
            WHERE hospital_id = %s AND year = %s AND scope = 'national'
            """,
            (hid, YEAR),
        )
        p = cur.fetchone()
        if not p:
            tags = [tag] if tag else []
            cur.execute(
                """
                INSERT INTO eqa_participation (hospital_id, year, scope, tags)
                VALUES (%s, %s, 'national', %s)
                RETURNING id
                """,
                (hid, YEAR, tags),
            )
            pid = cur.fetchone()["id"]
            stats.participations_created += 1
            if tag:
                stats.participations_tags_filled += 1
        else:
            pid = p["id"]
            stats.participations_existing += 1
            existing_tags = list(p["tags"] or [])
            # Only fill tags when currently empty — never overwrite non-empty.
            if tag and len(existing_tags) == 0:
                cur.execute(
                    "UPDATE eqa_participation SET tags = %s, updated_at = NOW() WHERE id = %s",
                    ([tag], pid),
                )
                stats.participations_tags_filled += 1

        for code in ones:
            cur.execute(
                """
                INSERT INTO eqa_participation_items (participation_id, field_code, value)
                VALUES (%s, %s, TRUE)
                ON CONFLICT (participation_id, field_code) DO NOTHING
                """,
                (pid, code),
            )
            if cur.rowcount:
                stats.items_inserted += 1
            else:
                stats.items_skipped_existing += 1


def ones_from_guo(row: dict, label_to_code: dict[str, str], stats: Stats) -> list[str]:
    out = []
    for label, v in (row.get("subjects") or {}).items():
        if v == 1 or v == "1" or v is True:
            code = label_to_code.get(norm_label(label))
            if not code:
                stats.subject_unmapped.add(label)
                continue
            out.append(code)
    return out


def run(xlsx: Path, dsn: str, dry: bool, commit: bool) -> Stats:
    data = load_xlsx(xlsx)
    stats = Stats()

    with psycopg.connect(dsn) as conn:
        conn.autocommit = False
        label_to_code = load_field_maps(conn)
        by_code, by_name = load_db_hospitals(conn)

        # Map hospital key -> id as we go
        id_by_key: dict[str, int] = {}

        for key, row in data["hospitals"].items():
            existing = match_hospital(row, by_code, by_name)
            dual = parse_dual_major(row.get("dual_raw"))
            if existing:
                stats.hospitals_matched += 1
                hid = existing["id"]
                fill_empty_hospital(conn, hid, row, existing, stats, dry)
                ensure_dual_major(conn, hid, dual, existing.get("dual_major"), stats, dry)
                # refresh in-memory code index if we filled code
                nc = norm_code(row["code"]) or norm_code(existing.get("customer_code"))
                if nc:
                    by_code[nc] = {**existing, "id": hid, "customer_code": row["code"] or existing.get("customer_code")}
            else:
                hid = insert_hospital(conn, row, stats, dry)
                if hid is None:
                    stats.hospitals_skipped += 1
                    continue
                if dry:
                    ensure_dual_major(conn, -1, dual, None, stats, dry=True)
                elif hid:
                    ensure_dual_major(conn, hid, dual, None, stats, dry=False)
                if not dry and hid:
                    nc = norm_code(row["code"])
                    if nc:
                        by_code[nc] = {
                            "id": hid,
                            "name": row["name"],
                            "customer_code": row["code"],
                            "region": row.get("region") or "",
                            "branch_office": row.get("branch") or "",
                            "dual_major": dual,
                        }
            if hid and hid != -1:
                id_by_key[key] = hid
            elif dry:
                id_by_key[key] = -1

        # National participation from 国考客户
        for key, grow in data["guo_by_key"].items():
            hid = id_by_key.get(key)
            if hid is None:
                # matched earlier via hospitals universe
                base = data["hospitals"].get(key) or grow
                existing = match_hospital(base, by_code, by_name)
                if existing:
                    hid = existing["id"]
                else:
                    stats.errors.append(f"no hospital id for guo row {grow.get('code')} {grow.get('name')}")
                    continue
            ones = ones_from_guo(grow, label_to_code, stats)
            tag = grow.get("tag") or (data["hospitals"].get(key) or {}).get("tag")
            if ones or tag:
                ensure_national_participation(conn, hid, tag, ones, stats, dry)

        # Hospitals with tags but no national guo row / no ones: shell participation
        for key, row in data["hospitals"].items():
            if key in data["guo_by_key"]:
                continue
            tag = row.get("tag")
            if not tag:
                continue
            hid = id_by_key.get(key)
            if hid is None:
                continue
            # prefer national shell (per spec)
            ensure_national_participation(conn, hid, tag, [], stats, dry)

        if dry or not commit:
            conn.rollback()
            if dry:
                print("DRY-RUN rolled back")
            else:
                print("Parsed OK but --commit not set; rolled back")
        else:
            conn.commit()
            print("COMMITTED")

    return stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    root = Path(__file__).resolve().parents[1]
    ap.add_argument("--xlsx", type=Path, default=root / "scripts/data/EQA_v1.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--commit", action="store_true", help="Actually commit (default rolls back unless set)")
    ap.add_argument("--dsn", default=os.environ.get("DATABASE_URL", ""))
    args = ap.parse_args()
    if not args.dsn:
        print("DATABASE_URL or --dsn required", file=sys.stderr)
        sys.exit(2)
    if not args.xlsx.exists():
        print("xlsx missing:", args.xlsx, file=sys.stderr)
        sys.exit(2)
    # dry-run implies no commit
    commit = bool(args.commit) and not args.dry_run
    stats = run(args.xlsx, args.dsn, dry=args.dry_run, commit=commit)
    import json
    print(json.dumps(stats.as_dict(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
